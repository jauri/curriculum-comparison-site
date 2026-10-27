"""Parse Georgia Tech's general 8-semester CS plan into one row per course or slot.

Georgia Tech's school-specific plan parser (step 13, batch 3). Reads the newest saved
College of Computing credit sheet (data/raw/gatech/*_8-semester-plan-25-26.xlsx,
"General 8 Semester Plan, No Threads", AY 2025-26) and writes
data/processed/gatech/plan_parsed.csv with the columns in common.PARSED_COLUMNS.

The sheet is an Excel workbook read with openpyxl (data_only=True reads the values
Excel last computed for formula cells, not the formulas). On "Sheet1":
- A year label ("FIRST YEAR" ... "FOURTH YEAR") sits in column A of a header row whose
  columns B and F say "Fall" and "Spring";
- below it, each row has a Fall entry in column B with its hours in C, and a Spring
  entry in F with its hours in G; either side may be empty;
- a "TOTAL SEMESTER HOURS" row closes each year with the two semester totals;
- E4 states "126 Hours for Degree"; rows 57-63 hold the degree notes (Note 1-3).
The second sheet ("Advisor info") is an advising aid about social science choices and
is not read.

Entries are free text written by advisors, with abbreviated titles ("ENGL 1101 Engl
Comp I", "MATH 1551 "), so course titles come from the 2026-2027 catalog's course pages
(parse_gatech_descriptions.py's input) and the sheet's own text is kept in notes.

Modeling decisions (step 13, batch 3):
- The degree is built from two Threads; this is the plan before threads are chosen.
  Courses the sheet names are `required` (the common core; specs/05-comparison-set.md),
  even when it marks them "(Thread Hours)" (the CS core counts toward both threads),
  and each "Thread Hours" entry is an elective_slot for a thread course.
- "HIST 2111 or HIST 2112", "CS 2050 or 2051", and "APPH 1040 or 1050 or 1060" are
  choices: one elective_slot row with the option codes in `options`.
- The ethics course (Note 3) and the probability and statistics option (Note 2) list
  their courses, so those are the slots' `options` (SLOT_OPTIONS; the ISYE 2027/2028
  pair is listed by its first course and CEE/ISYE 3770 by ISYE 3770).
- The other entries are slots: Lab Science (Note 1), the social science and humanities
  electives, the ethics course (Note 3), the probability and statistics option
  (Note 2), and Junior Design (Part 1) and (Part 2), the two-semester capstone design
  sequence, whose course options the sheet doesn't list.
- The catalog states "124 credit hours for graduation, plus a two-hour Wellness
  course", which is the sheet's 126.

Usage:
    python src/parse_gatech.py
"""

import re

import openpyxl
import pandas as pd
from bs4 import BeautifulSoup

from common import PARSED_COLUMNS, finish_choices, load_school, rel

SCHOOL = load_school("gatech")
OUT_CSV = SCHOOL.processed / "plan_parsed.csv"
PLAN = "General 8 Semester Plan, No Threads, AY 2025-26 (College of Computing credit sheet)"

YEARS = {"FIRST YEAR": 1, "SECOND YEAR": 2, "THIRD YEAR": 3, "FOURTH YEAR": 4}
# (entry column, hours column, term) for the two halves of a row.
SIDES = [("B", "C", "Fall"), ("F", "G", "Spring")]
# A Georgia Tech course code at the start of an entry: "CS 1301", "ENGL 1101".
CODE = re.compile(r"^([A-Z]{2,4}) (\d{4})\b")
# "CS 2050 or 2051" / "APPH 1040 or 1050 or 1060": a code, then " or " + more numbers.
OR_NUMBERS = re.compile(r"^([A-Z]{2,4}) (\d{4})((?: or \d{4})+)\b")
# "HIST 2111 or HIST 2112": full codes joined by " or ".
OR_CODES = re.compile(r"\b([A-Z]{2,4} \d{4}) or ([A-Z]{2,4} \d{4})\b")

# Courses whose 2026-27 catalog title differs from the 2025-26 sheet's name for them
# (found by parse_gatech_descriptions.py's title listing).
RENAMED = {"CS 1100": " The 2025-26 sheet calls it 'Freshman Leap'; the 2026-27 catalog title is 'Pathways in "
                      "Computing'."}
NOTE_1 = "One lab science must be PHYS 2211, and two lab sciences must be in the same discipline (Note 1)."
# Slot entries (lowercased, spaces collapsed, starting text) -> (title, type, notes).
SLOTS = [
    ("lab science", ("Lab Science", "elective_slot", "Georgia Tech: three lab science courses. " + NOTE_1)),
    ("social science elective: suggested fulfillment: us constitution",
     ("Social Science Elective (US Constitution)", "gen_ed",
      "Suggested: POL 1101, INTA 1200, or PUBP 3000 (not both POL 1101 and INTA 1200), which also satisfies "
      "Georgia's US/GA Constitution requirement, to be met before 90 credit hours.")),
    ("social science elective/psyc 1101", ("Social Science Elective", "gen_ed",
      "The sheet suggests PSYC 1101; any approved social science course.")),
    ("social science elective", ("Social Science Elective", "gen_ed", "Georgia Tech: an approved social science course.")),
    ("humanities elective", ("Humanities Elective", "gen_ed", "Georgia Tech: an approved humanities course.")),
    ("ethics course", ("Ethics Course", "elective_slot",
      "One of CS 3001, CS 4001, CS 4002, CS 4003, CS 4726, or SLS 3110 (Note 3).")),
    ("probability & statistics option", ("Probability & Statistics Option", "elective_slot",
      "MATH 3215, MATH 3670, CEE/ISYE 3770, or ISYE 2027 and ISYE 2028 (Note 2; the ISYE pair uses 3 hours of "
      "free electives).")),
    ("junior design (part 1)", ("Junior Design (Part 1)", "elective_slot",
      "The first semester of the College of Computing's two-semester Junior Design capstone sequence; the "
      "sheet doesn't list the course options.")),
    ("junior design (part 2)", ("Junior Design (Part 2)", "elective_slot",
      "The second semester of the Junior Design capstone sequence.")),
    ("thread hours", ("Thread Course", "elective_slot",
      "A course required by or chosen from one of the student's two Threads (the degree is built from two of "
      "eight Threads).")),
]


# Options of the slots whose notes list them (Note 2, Note 3).
SLOT_OPTIONS = {
    "Ethics Course": ["CS 3001", "CS 4001", "CS 4002", "CS 4003", "CS 4726", "SLS 3110"],
    "Probability & Statistics Option": ["MATH 3215", "MATH 3670", "ISYE 3770", "ISYE 2027"],
}


# Lowercase and collapse whitespace (the sheet has double spaces and line breaks).
def norm(s: str) -> str:
    """Return s lowercased with whitespace collapsed."""
    return re.sub(r"\s+", " ", s).strip().lower()


# Catalog titles for every course on the saved 2026-27 subject pages, used because the
# sheet abbreviates titles. Read straight from the course-block title lines.
def catalog_titles() -> dict:
    """Return {code: title} from data/raw/gatech/*_courses-*.html."""
    titles = {}
    for path in sorted(SCHOOL.raw.glob("*_courses-*.html")):
        soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
        for p in soup.select("p.courseblocktitle"):
            t = re.sub(r"\s+", " ", p.get_text(" ").replace("\xa0", " ")).strip()
            # The title runs to the ". N Credit Hours" ending, so a period inside it
            # ("Survey of U.S. History I") doesn't cut it short.
            m = re.match(r"([A-Z]{2,4} \d{4})\. (.*?)\. [\d-]+(?: or [\d-]+)? Credit Hours?", t)
            if m:
                titles[m.group(1)] = m.group(2)
    return titles


# Turn one sheet entry and its hours into a row.
def parse_entry(entry: str, hours, titles: dict) -> dict:
    """Return the plan row for one cell of the credit sheet."""
    raw = re.sub(r"\s+", " ", entry).strip()
    row = {"course_code": "", "subject": "", "number": "", "title": "", "credits_min": int(hours),
           "credits_max": int(hours), "requirement_type": "required", "gened_tag": "", "min_grade": "",
           "counts_toward_total": True, "notes": "", "options": ""}
    thread = " Counted as Thread Hours on the sheet." if "(Thread Hours)" in raw else ""
    key = norm(raw)
    # Slots first: a slot entry may name suggested courses ("POL 1101 or INTA 1200") that
    # are not a choice the plan requires. next() with a default: the first slot whose key
    # starts the entry (SLOTS lists the longer keys first), or None.
    slot = next((v for k, v in SLOTS if key.startswith(k)), None)
    if slot is not None:
        title, rtype, notes = slot
        row.update(title=title, requirement_type=rtype, notes=notes, options="; ".join(SLOT_OPTIONS.get(title, [])))
        return row
    m_nums, m_codes = OR_NUMBERS.match(raw), OR_CODES.search(raw)
    if m_nums or m_codes:
        if m_nums:
            subj = m_nums.group(1)
            codes = [f"{subj} {m_nums.group(2)}"] + [f"{subj} {n}" for n in re.findall(r"\d{4}", m_nums.group(3))]
        else:
            codes = list(m_codes.groups())
        subjects = {c.split()[0] for c in codes}
        row.update(subject=subjects.pop() if len(subjects) == 1 else "",
                   title=f"{' or '.join(codes)} ({titles[codes[0]]})", requirement_type="elective_slot",
                   options="; ".join(codes),
                   notes=f"Student chooses one: {' | '.join(titles[c] for c in codes)}. Sheet: '{raw}'.{thread}")
        return row
    m = CODE.match(raw)
    if m:
        code = m.group(0)
        row.update(course_code=code, subject=m.group(1), number=m.group(2), title=titles[code],
                   notes=f"Sheet: '{raw}'.{thread}{RENAMED.get(code, '')}")
        return row
    raise SystemExit(f"Unknown entry on the credit sheet: {raw!r}. Add it to SLOTS.")


# Walk Sheet1 from top to bottom. A year label in column A starts a year; then each
# row's two sides become rows until the "TOTAL SEMESTER HOURS" row, whose C and G
# cells are the stated semester totals.
def parse_plan(ws, titles: dict) -> tuple[pd.DataFrame, dict]:
    """Return (rows, {(year, term): stated hours})."""
    rows, term_hours, year = [], {}, None
    for r in range(1, ws.max_row + 1):
        label = ws[f"A{r}"].value
        if isinstance(label, str) and label.strip() in YEARS:
            year = YEARS[label.strip()]
            continue
        if year is None:
            continue  # the header rows above the first year
        if ws[f"B{r}"].value == "TOTAL SEMESTER HOURS":
            term_hours[(year, "Fall")], term_hours[(year, "Spring")] = ws[f"C{r}"].value, ws[f"G{r}"].value
            if year == 4:
                break  # the notes and student-information rows follow
            continue
        for entry_col, hours_col, term in SIDES:
            entry, hours = ws[f"{entry_col}{r}"].value, ws[f"{hours_col}{r}"].value
            if entry in (None, "") or not str(entry).strip():
                continue
            row = parse_entry(str(entry), hours, titles)
            row.update(orig_year=year, orig_term=term)
            rows.append(row)
    df = pd.DataFrame(rows)
    df["orig_calendar_year"] = ""
    df["orig_term_type"] = "academic"
    df["is_coop"] = False
    return df, term_hours


# Run the parse checks and return a list of failure messages (empty = all passed).
def check(df: pd.DataFrame, term_hours: dict, stated: int, catalog_text: str) -> list[str]:
    """Return failed-check messages."""
    failures = []
    total = df["credits_min"].sum()
    if total != stated:
        failures.append(f"Credits {total} != the sheet's {stated} Hours for Degree.")
    if SCHOOL.credits != (stated, stated):
        failures.append(f"schools.csv says {SCHOOL.credits} credits but the sheet states {stated}.")
    if "a total of 124 credit hours for graduation, plus a two-hour Wellness course" not in catalog_text:
        failures.append("The catalog no longer states 124 hours plus a two-hour Wellness course; recheck 126.")
    for (y, t), hours in term_hours.items():
        got = df[(df["orig_year"] == y) & (df["orig_term"] == t)]["credits_min"].sum()
        if got != hours:
            failures.append(f"Year {y} {t}: {got} hours, sheet says {hours}.")
    if len(term_hours) != 8:
        failures.append(f"Expected 8 semesters, found {len(term_hours)}.")
    return failures


# Script entry point: parse, add the program columns, check, and write the CSV.
def main() -> None:
    path = SCHOOL.newest(SCHOOL.credits_source)
    wb = openpyxl.load_workbook(path, data_only=True)  # data_only: computed values, not formulas
    ws = wb["Sheet1"]
    if ws["A3"].value != "Thread Plan - AY 2025-26":
        raise SystemExit(f"{rel(path)}: A3 is {ws['A3'].value!r}, expected the AY 2025-26 thread plan.")
    m = re.match(r"(\d+) Hours for Degree", str(ws["E4"].value))
    if not m:
        raise SystemExit(f"{rel(path)}: E4 does not state the hours for the degree.")
    catalog = SCHOOL.newest(SCHOOL.college_source)
    catalog_text = re.sub(r"\s+", " ", BeautifulSoup(catalog.read_text(encoding="utf-8"), "html.parser")
                          .get_text(" ").replace("\xa0", " "))
    if SCHOOL.college not in catalog_text:
        raise SystemExit(f"College '{SCHOOL.college}' not found on {rel(catalog)}.")
    df, term_hours = parse_plan(ws, catalog_titles())

    df["institution"] = SCHOOL.institution
    df["program"] = SCHOOL.program
    df["catalog_year"] = SCHOOL.catalog_year
    df["plan"] = PLAN
    df["source_file"] = rel(path)

    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    finish_choices(df)[PARSED_COLUMNS].to_csv(OUT_CSV, index=False)  # choice rules: common.py
    print(f"Wrote {len(df)} rows to {rel(OUT_CSV)}")
    print(f"Credits: {df['credits_min'].sum()} (the sheet states {m.group(1)})")
    failures = check(df, term_hours, int(m.group(1)), catalog_text)
    if failures:
        raise SystemExit("CHECKS FAILED:\n  " + "\n  ".join(failures))
    print("All step 13 parse checks passed.")


# Run main() only when executed as a script, not when imported.
if __name__ == "__main__":
    main()
